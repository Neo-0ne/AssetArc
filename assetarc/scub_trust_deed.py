import re
import os
from docx import Document

def scrub_document(input_path, output_path):
    """
    Reads a .docx file, replaces personal information with placeholders,
    and saves the modified document.
    """
    try:
        doc = Document(input_path)
        for paragraph in doc.paragraphs:
            # Replace names
            paragraph.text = re.sub(r'JOHANNES JACOBUS BEKKER', '<<FOUNDER_NAME>>', paragraph.text)
            paragraph.text = re.sub(r'NICOLAAS ESAIAS DE VILLIERS', '<<TRUSTEE_1_NAME>>', paragraph.text)
            paragraph.text = re.sub(r'RENATE SYPKENS', '<<TRUSTEE_2_NAME>>', paragraph.text)

            # Replace ID numbers
            paragraph.text = re.sub(r'831226 5017 088', '<<FOUNDER_ID>>', paragraph.text)
            paragraph.text = re.sub(r'760219 5005 083', '<<TRUSTEE_1_ID>>', paragraph.text)
            paragraph.text = re.sub(r'850819 0008 080', '<<TRUSTEE_2_ID>>', paragraph.text)

            # Replace addresses
            paragraph.text = re.sub(r'5 Gordon Smit Screscent, Dan Pienaar, Bloemfontein, South Africa', '<<FOUNDER_ADDRESS>>', paragraph.text)
            paragraph.text = re.sub(r'Nr. 7, JD Potgieter Street, Tempe, Bloemfontein, South Africa', '<<TRUSTEE_1_ADDRESS>>', paragraph.text)
            paragraph.text = re.sub(r'24 McArthur Street, Fichardtpark, Bloemfontein, South Africa', '<<TRUSTEE_2_ADDRESS>>', paragraph.text)

        doc.save(output_path)
        print(f"Scrubbed document saved to {output_path}")

    except Exception as e:
        print(f"An error occurred: {e}")

if __name__ == '__main__':
    input_file = 'assetarc/bots/Trust_Bot/templates/trust_template.docx'
    output_file = 'assetarc/bots/Trust_Bot/templates/trust_template_scrubbed.docx'
    scrub_document(input_file, output_file)
